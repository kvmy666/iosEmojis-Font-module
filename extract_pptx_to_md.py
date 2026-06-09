"""
Extract all slide text (titles, body, tables) from the 8 Git (FWDN 232) PPTX
lectures into a single Markdown file: git_full_notes.md
No manual reading of the slides — pure python-pptx extraction.
"""
import os
from pptx import Presentation

PPTX_DIR = "/root/.claude/uploads/ca553555-cd29-5783-b4d9-921ffa4bffcf/"
PPTX = [
    ("b25f87a9-wdt_01.pptx", "Chapter 1: Introduction to Git & Version Control"),
    ("039d2739-wdt_02.pptx", "Chapter 2: Getting Started with Git"),
    ("ac754add-wdt_03.pptx", "Chapter 3: Branching in Git"),
    ("464f8d7a-wdt_05.pptx", "Chapter 5: Git Workflows"),
    ("9e875f97-wdt_06.pptx", "Chapter 6: Correcting Errors While Working with Git"),
    ("a7d1d0e5-wdt_07.pptx", "Chapter 7: Unlocking Git's Full Potential"),
    ("a8503aab-wdt_08.pptx", "Chapter 8: Integrate Git in Your Development Cycle"),
    ("44ac2990-wdt_09.pptx", "Chapter 9: Git GUI Tools"),
]


def shape_text(shape):
    out = []
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            line = "".join(run.text for run in para.runs).strip()
            if line:
                out.append(line)
    if shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                out.append(" | ".join(cells))
    return out


def main():
    output = ["# Git / Web Developer Tools (FWDN 232) — Full Lecture Notes\n"]
    for filename, title in PPTX:
        path = os.path.join(PPTX_DIR, filename)
        output.append(f"\n# {title}\n")
        prs = Presentation(path)
        for i, slide in enumerate(prs.slides, 1):
            lines = []
            for shape in slide.shapes:
                lines.extend(shape_text(shape))
            if lines:
                output.append(f"\n## Slide {i}\n")
                output.append("\n".join(lines))
                output.append("\n---")
    with open("git_full_notes.md", "w", encoding="utf-8") as f:
        f.write("\n".join(output))
    print("Saved: git_full_notes.md")


if __name__ == "__main__":
    main()
